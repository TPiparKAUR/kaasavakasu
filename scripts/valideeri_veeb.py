"""Veebilehe ja Exceli kuvatud väärtuste kontroll (V8–V10): võrdleb brauseris kuvatud numbreid tulemusfailidega.

Metoodika v2: kuvatav põhiväärtus on mediaan; stsenaariumi koguhulk pindala kohta tuleb keskmiste vahest.

V8  Tulemustabel (Tüübid x seisund): iga kuvatud lahter (mediaan) vs ``kaasnev_tuup_seisund.parquet``
V9  Stsenaarium: kuvatud vilets/hea mediaanid, mediaanide vahe, keskmiste vahe ja keskmiste vahe x pindala
V10 Excel ``Tüübid × seisund`` (mediaan) vs ``kaasnev_tuup_seisund.parquet``

Käivitus: ``uv run --with playwright --with openpyxl python scripts/valideeri_veeb.py`` (openpyxl ja playwright
ei ole projekti sõltuvused; kasutatakse ainult kontrolliks).
"""

from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path

import polars as pl
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
SITE = ROOT / "site"
RES = ROOT / "data" / "tulemused"


def num(t: str) -> float | None:
    s = re.sub(r"[\s  ]", "", t).replace("−", "-").replace(",", ".").replace("+", "")
    try:
        return float(s)
    except ValueError:
        return None


def close(shown: float, true: float, tol: float = 0.007) -> bool:
    return abs(shown - true) <= max(tol * abs(true), 6e-4)


def near(shown: float | None, true: float, floor: float) -> bool:
    """Kuvatud vahe (ümardatud) vs täpne vahe: 1,2 % või absoluutne alampiir."""
    return shown is not None and abs(shown - true) <= max(0.012 * abs(true), floor)


def expected_rows(kt: pl.DataFrame, layers: list[str]) -> list[dict]:
    rows = []
    d = kt.filter(pl.col("okosusteem_kood").is_in([1, 2, 3, 4, 5]) & pl.col("uld").is_in([1, 2, 3]))
    for _key, g in sorted(d.group_by("okosusteem_kood", "alam_kood", "uld"), key=lambda kv: kv[0]):
        by = {r["layer_id"]: r for r in g.iter_rows(named=True)}
        r0 = g.row(0, named=True)
        rows.append({"tuup": r0["tuup"], "seisund": r0["seisund_nimi"], "area": r0["pindala_ha"], "cells": [by.get(lid) for lid in layers]})
    return rows


def main() -> None:
    kt = pl.read_parquet(RES / "kaasnev_tuup_seisund.parquet")
    K = json.loads((SITE / "data/kaasnev.json").read_text(encoding="utf-8"))
    layers = [x["id"] for x in K["kihid"]]
    exp = expected_rows(kt, layers)
    out: dict = {}
    with sync_playwright() as p:
        exe = os.environ.get("CHROMIUM_PATH")  # vaikimisi Playwrighti oma brauser
        b = p.chromium.launch(executable_path=exe, args=["--no-sandbox"])
        pg = b.new_page(viewport={"width": 1400, "height": 900})
        errs: list[str] = []
        pg.on("pageerror", lambda e: errs.append(str(e)))
        pg.goto(f"file://{SITE}/oks-seisund.html")
        pg.wait_for_selector("section.hero")
        pg.click("nav.sections a:has-text('Tulemustabel')")
        pg.wait_for_selector("table.rt tr.dr")
        shown = pg.evaluate("""() => [...document.querySelectorAll('table.rt tr.dr')].map(tr => ({tuup: tr.querySelector('th.fx0').textContent, seisund: tr.querySelector('td.fx1').textContent, area: tr.querySelector('td.fx2').textContent, cells: [...tr.querySelectorAll('td[data-r]')].map(td => td.textContent)}))""")
        bad, n, textual = [], 0, 0
        if len(shown) != len(exp):
            out["V8 ridade arv"] = {"veeb": len(shown), "oodatud": len(exp)}
        for s_row, e_row in zip(shown, exp, strict=False):
            if s_row["tuup"] != e_row["tuup"] or s_row["seisund"].strip() != e_row["seisund"]:
                bad.append(("rida", s_row["tuup"], e_row["tuup"]))
                continue
            if abs((num(s_row["area"]) or 0) - e_row["area"]) > 0.6:
                bad.append(("pindala", s_row["tuup"], s_row["area"], e_row["area"]))
            for lid, txt, e in zip(layers, s_row["cells"], e_row["cells"], strict=True):
                if e is None:
                    continue
                st = e["staatus"]
                if st == "ei paku":
                    textual += 1
                    if "ei paku" not in txt:
                        bad.append((e_row["tuup"], e_row["seisund"], lid, txt, st))
                elif st == "ei hinnatud":
                    textual += 1
                    if txt.strip() != "–":
                        bad.append((e_row["tuup"], e_row["seisund"], lid, txt, st))
                else:
                    v = num(txt)
                    n += 1
                    if v is None or not close(v, e["median"]):
                        bad.append((e_row["tuup"], e_row["seisund"], lid, txt, e["median"]))
        out["V8 tulemustabel"] = {"ridu veebis": len(shown), "arvulisi lahtreid": n, "tekstilisi (ei paku / –)": textual, "lahknevusi": len(bad), "näited": bad[:5]}
        # V9 stsenaarium: Salumets vilets -> hea, 1000 ha
        pg.click("nav.sections a:has-text('Stsenaarium')")
        pg.wait_for_selector("table.scn")
        unit = pg.evaluate("document.querySelector('#cu').selectedOptions[0].textContent")
        ha = float(pg.input_value("#ch"))
        frm, to = pg.input_value("#cf"), pg.input_value("#ct")
        rows = pg.evaluate("""() => [...document.querySelectorAll('table.scn tbody tr')].map(tr => [...tr.querySelectorAll('td')].map(td => td.textContent))""")
        names = {x["id"]: (x["nimi"], x["roll"]) for x in K["kihid"]}
        unit = unit.strip()
        sel = kt.filter(pl.col("tuup") == unit)
        uld = {"hea": 1, "keskmine": 2, "vilets": 3}
        bad9, n9 = [], 0
        by_nm: dict[str, list] = {}
        for r in rows:
            by_nm.setdefault(r[0], []).append(r)
        for lid in layers:
            nm, roll = names[lid]
            rr = [x for x in rows if x[0].startswith(nm) and x[0].endswith({"bio": "biofüüsikaline", "raha": "rahaline", "lisa": "lisakiht"}[roll])]
            a = sel.filter((pl.col("layer_id") == lid) & (pl.col("uld") == uld[frm]))
            c = sel.filter((pl.col("layer_id") == lid) & (pl.col("uld") == uld[to]))
            if not rr or a.is_empty() or c.is_empty() or a["mean"][0] is None or c["mean"][0] is None or a["staatus"][0] in ("ei paku", "ei hinnatud") or c["staatus"][0] in ("ei paku", "ei hinnatud"):
                continue
            row = rr[0]
            # veerud: nimi, ühik, mediaan a, mediaan b, mediaanide vahe, keskmiste vahe, keskmiste vahe x pindala
            va, vb, vdm, vd = num(row[2]), num(row[3]), num(row[4]), num(row[5])
            vt = num(row[6].split(" ")[0] if row[6] else "")
            ma, mb = a["median"][0], c["median"][0]
            ta, tb = a["mean"][0], c["mean"][0]
            n9 += 1
            floor = 1e-4 * max(abs(ta), abs(tb), abs(ma), abs(mb)) + 6e-4
            ok = va is not None and vb is not None and close(va, ma) and close(vb, mb)
            ok = ok and near(vdm, mb - ma, floor) and near(vd, tb - ta, floor)
            if "/ha" in row[1] and "0,01" not in row[1] and vt is not None:
                ok = ok and abs(vt - (tb - ta) * ha) <= max(0.012 * abs((tb - ta) * ha), 0.6)
            if not ok:
                bad9.append((nm, roll, row[2:7], ma, mb, ta, tb))
        out["V9 stsenaarium"] = {"üksus": unit, "seisundid": f"{frm}→{to}", "pindala ha": ha, "kontrollitud ridu": n9, "lahknevusi": len(bad9), "näited": bad9[:4]}
        out["JS vead brauseris"] = errs
        b.close()
    # V10 Excel
    from openpyxl import load_workbook

    wb = load_workbook(SITE / "andmetabel_kaasneva_kasu.xlsx", read_only=True, data_only=True)
    ws = wb["Tüübid × seisund"]
    data = list(ws.iter_rows(min_row=4, values_only=True))
    ix = {(r["tuup"], r["seisund_nimi"], r["layer_id"]): r for r in kt.filter(pl.col("uld").is_in([1, 2, 3, 4, 5])).iter_rows(named=True)}
    bad10 = n10 = 0
    for row in data:
        if row[1] is None:
            continue
        for lid, v in zip(layers, row[4:], strict=False):
            e = ix.get((row[1], row[2], lid))
            if e is None or v is None or isinstance(v, str) or e["median"] is None:
                continue
            n10 += 1
            if abs(v - e["median"]) > 5e-5 * max(abs(e["median"]), 1e-9):
                bad10 += 1
    out["V10 Excel (Tüübid × seisund)"] = {"ridu": sum(1 for r in data if r[1]), "arvulisi lahtreid": n10, "lahknevusi": bad10, "lehed": wb.sheetnames}
    print(json.dumps(out, ensure_ascii=False, indent=2, default=str))
    (RES / "validatsioon_veeb.json").write_text(json.dumps(out, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


if __name__ == "__main__":
    sys.exit(main())
